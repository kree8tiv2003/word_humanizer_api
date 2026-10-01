"""Turn uploads and links into a SourceDoc made of numbered units."""
from __future__ import annotations

import io
import ipaddress
import os
import re
import socket
from urllib.parse import urljoin, urlparse

from .models import SourceDoc, Unit

AUDIO_EXT = {'.mp3', '.wav', '.m4a', '.aac', '.flac', '.ogg', '.oga', '.opus', '.webm', '.mp4', '.wma', '.aiff', '.aif'}
TEXT_EXT = {'.txt', '.md', '.markdown', '.fountain', '.rtf', '.srt', '.lrc'}
MAX_UNIT_CHARS = 1400
MAX_FETCH_BYTES = int(os.getenv('MAX_FETCH_MB', '50')) * 1024 * 1024


class IngestError(ValueError):
    pass


def ext_of(name: str) -> str:
    return os.path.splitext(name or '')[1].lower()


def is_audio(name: str, content_type: str = '') -> bool:
    return ext_of(name) in AUDIO_EXT or content_type.startswith('audio/')


# --------------------------------------------------------------------------- text -> units
def _split_long(par: str) -> list[str]:
    if len(par) <= MAX_UNIT_CHARS:
        return [par]
    sentences = re.split(r'(?<=[.!?…])\s+', par)
    out, cur = [], ''
    for s in sentences:
        if cur and len(cur) + len(s) > MAX_UNIT_CHARS * 0.6:
            out.append(cur.strip()); cur = ''
        cur += s + ' '
    if cur.strip():
        out.append(cur.strip())
    return out


def text_to_units(text: str) -> list[Unit]:
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[ \t]+', ' ', text)
    blocks = [b.strip() for b in re.split(r'\n\s*\n', text) if b.strip()]
    if len(blocks) <= 2 and text.count('\n') > 6:      # line-based text such as lyrics
        blocks = [b.strip() for b in text.split('\n') if b.strip()]
    pars = []
    for b in blocks:
        lines = b.split('\n')
        # Keep short line structure (lyrics, poetry, dialogue); join wrapped prose.
        if len(lines) > 1 and sum(len(l) for l in lines) / len(lines) > 60:
            b = ' '.join(l.strip() for l in lines)
        pars.extend(_split_long(b))
    return [Unit(idx=i, text=p) for i, p in enumerate(pars)]


def _strip_rtf(raw: str) -> str:
    raw = re.sub(r'\\par[d]?', '\n', raw)
    raw = re.sub(r'\{\\\*[^{}]*\}', '', raw)
    raw = re.sub(r'\\[a-zA-Z]+-?\d* ?', '', raw)
    return raw.replace('{', '').replace('}', '')


def _strip_subtitles(raw: str) -> str:
    raw = re.sub(r'^\d+\s*$', '', raw, flags=re.M)                       # SRT counters
    raw = re.sub(r'^[\d:,.]+\s*-->\s*[\d:,.]+.*$', '', raw, flags=re.M)  # SRT timings
    raw = re.sub(r'^\[\d+:\d+(?:\.\d+)?\]', '', raw, flags=re.M)        # LRC timings
    return re.sub(r'\n{2,}', '\n', raw)


def decode_text(data: bytes) -> str:
    for enc in ('utf-8-sig', 'utf-16', 'cp1252', 'latin-1'):
        try:
            s = data.decode(enc)
            if enc == 'utf-16' and not data[:2] in (b'\xff\xfe', b'\xfe\xff'):
                continue
            return s
        except UnicodeDecodeError:
            continue
    return data.decode('utf-8', errors='replace')


# --------------------------------------------------------------------------- formats
def from_pdf(data: bytes) -> str:
    import pymupdf
    try:
        doc = pymupdf.open(stream=data, filetype='pdf')
    except Exception as e:
        raise IngestError(f'Could not open the PDF: {e}')
    pars = []
    for page in doc:
        for block in page.get_text('dict')['blocks']:
            if block.get('type') != 0:   # image block
                continue
            lines = []
            for ln in block['lines']:
                txt = ''.join(s['text'] for s in ln['spans']).strip()
                if txt:
                    lines.append((txt, ln['bbox']))
            if not lines:
                continue
            left = min(b[0] for _, b in lines)
            right = max(b[2] for _, b in lines)
            heights = sorted(b[3] - b[1] for _, b in lines)
            lh = heights[len(heights) // 2] or 10
            cur = ''
            for i, (txt, bb) in enumerate(lines):
                cur = (cur[:-1] + txt) if cur.endswith('-') and txt[:1].islower() else (cur + ' ' + txt).strip()
                nxt = lines[i + 1][1] if i + 1 < len(lines) else None
                short_end = bb[2] < left + 0.82 * (right - left) and re.search(r'[.!?…"”’:)]$', txt)
                gap = nxt is not None and nxt[1] - bb[3] > 0.8 * lh
                indent = nxt is not None and nxt[0] - left > 1.5 * lh
                if nxt is None or short_end or gap or indent:
                    if cur and not re.fullmatch(r'\d{1,4}', cur):  # drop bare page numbers
                        pars.append(cur)
                    cur = ''
    if not pars:
        raise IngestError('The PDF has no extractable text (it may be scanned images).')
    return '\n\n'.join(pars)


def from_docx(data: bytes) -> str:
    import docx
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise IngestError(f'Could not open the Word document: {e}')
    pars = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                pars.append(' | '.join(dict.fromkeys(cells)))
    if not pars:
        raise IngestError('The Word document has no text.')
    return '\n\n'.join(pars)


def from_html(html: str, url: str = '') -> tuple[str, str]:
    title = ''
    m = re.search(r'<title[^>]*>(.*?)</title>', html, re.I | re.S)
    if m:
        title = re.sub(r'\s+', ' ', m.group(1)).strip()
    text = ''
    try:
        import trafilatura
        text = trafilatura.extract(html, url=url or None, include_comments=False, include_tables=False,
                                   favor_recall=True) or ''
    except Exception:
        text = ''
    if len(text) < 200:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, 'html.parser')
        for tag in soup(['script', 'style', 'nav', 'header', 'footer', 'aside', 'form', 'noscript']):
            tag.decompose()
        root = soup.find('article') or soup.find('main') or soup.body or soup
        pars = [p.get_text(' ', strip=True) for p in root.find_all(['p', 'h1', 'h2', 'h3', 'li', 'blockquote', 'pre'])]
        alt = '\n\n'.join(p for p in pars if p)
        if len(alt) > len(text):
            text = alt
    if not text.strip():
        raise IngestError('No readable story text was found on that page.')
    return title, text


def load_document(filename: str, data: bytes) -> SourceDoc:
    """Load a non-audio upload."""
    ext = ext_of(filename)
    title = os.path.splitext(os.path.basename(filename or 'Untitled'))[0]
    if ext == '.pdf' or data[:5] == b'%PDF-':
        text, kind = from_pdf(data), 'pdf'
    elif ext == '.docx' or (data[:2] == b'PK' and ext not in ('.zip',)):
        text, kind = from_docx(data), 'docx'
    elif ext == '.doc':
        raise IngestError('Old .doc files are not supported. Save it as .docx, PDF or text and upload again.')
    elif ext in ('.html', '.htm'):
        page_title, text = from_html(decode_text(data))
        title, kind = page_title or title, 'html'
    else:
        raw = decode_text(data)
        if ext == '.rtf' or raw.startswith('{\\rtf'):
            raw = _strip_rtf(raw)
        if ext in ('.srt', '.lrc', '.vtt'):
            raw = _strip_subtitles(raw)
        text, kind = raw, 'text'
    units = text_to_units(text)
    if not units:
        raise IngestError(f'{filename} contains no text.')
    return SourceDoc(kind=kind, title=title, units=units)


def merge_docs(docs: list[SourceDoc]) -> SourceDoc:
    """Combine several text uploads into one source, in upload order."""
    if len(docs) == 1:
        return docs[0]
    units, notes = [], []
    for d in docs:
        for u in d.units:
            units.append(Unit(idx=len(units), text=u.text))
        notes += d.notes
    return SourceDoc(kind='+'.join(sorted({d.kind for d in docs})), title=docs[0].title, units=units, notes=notes)


# --------------------------------------------------------------------------- links
def _check_host(url: str) -> None:
    p = urlparse(url)
    if p.scheme not in ('http', 'https') or not p.hostname:
        raise IngestError('Only http and https links are supported.')
    if os.getenv('ALLOW_PRIVATE_URLS') == '1':
        return
    try:
        infos = socket.getaddrinfo(p.hostname, p.port or (443 if p.scheme == 'https' else 80))
    except socket.gaierror:
        raise IngestError(f'Could not resolve {p.hostname}.')
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise IngestError('Links to private or local network addresses are blocked.')


def fetch_url(url: str) -> tuple[bytes, str, str]:
    """Download a link with redirect and size checks. Returns (data, content_type, final_url)."""
    import requests
    headers = {'User-Agent': 'Mozilla/5.0 (ScriptStudio; +script writer)'}
    for _ in range(6):
        _check_host(url)
        r = requests.get(url, headers=headers, timeout=25, stream=True, allow_redirects=False)
        if r.is_redirect or r.status_code in (301, 302, 303, 307, 308):
            url = urljoin(url, r.headers.get('location', ''))
            continue
        if r.status_code >= 400:
            raise IngestError(f'The link returned HTTP {r.status_code}.')
        buf = io.BytesIO()
        for chunk in r.iter_content(65536):
            buf.write(chunk)
            if buf.tell() > MAX_FETCH_BYTES:
                raise IngestError('The linked file is too large.')
        return buf.getvalue(), r.headers.get('content-type', '').split(';')[0].strip().lower(), url
    raise IngestError('Too many redirects.')


def guess_name(url: str, content_type: str) -> str:
    name = os.path.basename(urlparse(url).path) or 'link'
    if ext_of(name):
        return name
    for ct, ext in (('pdf', '.pdf'), ('wordprocessingml', '.docx'), ('html', '.html'), ('text/plain', '.txt'),
                    ('audio/mpeg', '.mp3'), ('audio/wav', '.wav'), ('audio/x-wav', '.wav'), ('audio/mp4', '.m4a'),
                    ('audio/ogg', '.ogg'), ('audio/flac', '.flac'), ('audio/webm', '.webm')):
        if ct in content_type:
            return name + ext
    return name + '.html'
